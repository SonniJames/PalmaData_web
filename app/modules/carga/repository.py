"""
PalmaData · Cargar datos · Repositorio
======================================
Lee un Excel y lo inserta en su tabla con ON CONFLICT DO NOTHING, igual
que hace el endpoint de la app cuando hay internet. Los triggers y las
funciones de cada tabla corren exactamente igual: a la base no le importa
por dónde llegó el dato.
"""
import hashlib
import re
from io import BytesIO

from openpyxl import load_workbook
from psycopg2 import errors
from psycopg2.extras import execute_values

from ...core import db
from .tablas import EXTRAS_CONOCIDAS, FIJOS, TABLAS, normalizar

# Cuántas filas por INSERT. Los archivos de tracks traen más de mil filas y
# mandarlas de una sola vez hace la instrucción enorme; en bloques el uso de
# memoria queda acotado y el tiempo es prácticamente el mismo.
LOTE = 500


def columnas_de_tabla(tabla: str) -> list[str]:
    """Las columnas reales de la tabla, leídas de la base, no supuestas."""
    filas = db.fetch_all("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = 'plantacion' AND table_name = %s
        ORDER BY ordinal_position
    """, (tabla,))
    return [f["column_name"] for f in filas]


# Un DEFAULT que sea un número suelto: 0, 0.0, -1. Se descartan a propósito
# los que son expresiones (nextval, now(), gen_random_uuid): esos los tiene
# que resolver PostgreSQL, no nosotros.
_DEFAULT_NUMERICO = re.compile(r"^-?\d+(\.\d+)?$")


def defaults_de_tabla(tabla: str) -> dict:
    """
    El valor por defecto numérico de cada columna que lo tenga.

    Hace falta por algo que no es evidente: un DEFAULT solo se aplica cuando
    la columna NO se menciona en el INSERT. Como el cargador las menciona
    todas, una celda vacía del Excel entraba como NULL y se saltaba el
    DEFAULT 0 de la tabla — mientras que por red el mismo dato llegaba en 0.
    El registro terminaba distinto según por dónde entrara.
    """
    filas = db.fetch_all("""
        SELECT column_name, column_default
        FROM information_schema.columns
        WHERE table_schema = 'plantacion' AND table_name = %s
          AND column_default IS NOT NULL
    """, (tabla,))
    salida = {}
    for f in filas:
        bruto = str(f["column_default"]).split("::")[0].strip().strip("'")
        if _DEFAULT_NUMERICO.match(bruto):
            salida[f["column_name"]] = float(bruto) if "." in bruto else int(bruto)
    return salida


def leer_excel(contenido: bytes) -> tuple[list[str], list[list]]:
    """
    Devuelve los encabezados y las filas con datos.

    Se descartan las filas totalmente vacías: Excel suele dejar cientos al
    final de la hoja y sin esto se intentarían insertar como registros en
    blanco.
    """
    wb = load_workbook(BytesIO(contenido), data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        filas_iter = ws.iter_rows(values_only=True)
        try:
            cabecera = next(filas_iter)
        except StopIteration:
            return [], []
        encabezados = [str(c).strip() if c is not None else "" for c in cabecera]
        datos = [list(f) for f in filas_iter if any(v is not None for v in f)]
        return encabezados, datos
    finally:
        wb.close()


def mapear(encabezados: list[str], columnas_bd: list[str]) -> tuple[dict, list[str]]:
    """
    Empareja cada encabezado del Excel con su columna en la base.

    Primero por nombre exacto; si no, comparando sin guiones bajos ni
    mayúsculas (ver `normalizar`). Las que no correspondan a ninguna columna
    se devuelven aparte para avisarlo: no se inventan ni se silencian.
    """
    por_normal = {normalizar(c): c for c in columnas_bd}
    mapa: dict[int, str] = {}
    desconocidas: list[str] = []

    for i, enc in enumerate(encabezados):
        if not enc:
            continue
        if enc in columnas_bd:
            mapa[i] = enc
        elif normalizar(enc) in por_normal:
            mapa[i] = por_normal[normalizar(enc)]
        elif normalizar(enc) in {normalizar(x) for x in EXTRAS_CONOCIDAS}:
            # El Excel sale de SQLite y trae alguna columna que la tabla no
            # tiene (`equipo`). server.py tampoco la manda por red, así que
            # no es una pérdida: se omite sin ensuciar el informe.
            continue
        else:
            desconocidas.append(enc)
    return mapa, desconocidas


def huella(contenido: bytes) -> str:
    """SHA-256 del archivo, para reconocerlo aunque lo renombren."""
    return hashlib.sha256(contenido).hexdigest()


def carga_previa(hash_archivo: str) -> dict | None:
    """Si este mismo contenido ya se cargó antes, devuelve aquella carga."""
    return db.fetch_one("""
        SELECT archivo, tabla, filas_nuevas, cargado_por, cargado_at
        FROM plantacion.carga_archivos
        WHERE huella = %s AND resultado = 'ok'
        ORDER BY cargado_at DESC LIMIT 1
    """, (hash_archivo,))


def ya_en_otra_tabla(tabla_otra: str, columna: str, valores: list) -> set:
    """
    Cuáles de estos identificadores ya están en la otra tabla.

    Es el caso de los tracks: entran a tracksmoviltemp y un cron los pasa a
    tracksmovil, vaciando la temporal. Si el archivo se vuelve a subir
    DESPUÉS de que el cron corrió, la temporal está vacía y el ON CONFLICT
    no vería nada: el registro entraría por segunda vez.
    """
    limpios = [v for v in valores if v is not None]
    if not limpios:
        return set()
    filas = db.fetch_all(
        f'SELECT "{columna}" AS v FROM plantacion.{tabla_otra} '
        f'WHERE "{columna}" = ANY(%s)', (limpios,))
    return {f["v"] for f in filas}


def insertar(tabla: str, mapa: dict[int, str], filas: list[list],
             columnas_bd: list[str]) -> tuple[int, int]:
    """
    Inserta las filas y devuelve (nuevas, repetidas).

    Las nuevas se cuentan con RETURNING: con ON CONFLICT el rowcount no es
    de fiar, y saber cuántas entraron de verdad es justo el dato que
    interesa.
    """
    conf = TABLAS[tabla]
    indices = sorted(mapa)
    columnas = [mapa[i] for i in indices]

    # Las columnas que server.py escribe con un valor fijo y el Excel no
    # trae. Van al final, con el mismo valor para todas las filas, para que
    # la fila quede idéntica a la que deja la carga por red.
    fijos = {c: v for c, v in FIJOS.get(tabla, {}).items()
             if c not in columnas and c in columnas_bd}

    # Para las celdas vacías: el DEFAULT de la columna, si es un número.
    # Sin esto, una casilla en blanco del Excel entra como NULL y se salta
    # el DEFAULT 0 que la tabla declara, cosa que por red no pasa.
    por_defecto = defaults_de_tabla(tabla)
    rellenos = [por_defecto.get(c) for c in columnas]
    destino = ", ".join(f'"{c}"' for c in columnas + list(fijos))

    # Si el registro pudo haber sido movido ya a otra tabla, se descartan
    # aquí, en Python, y no dentro del SQL. Hacerlo con
    # `SELECT * FROM (VALUES ...)` obligaba a PostgreSQL a adivinar los
    # tipos, los tomaba todos como texto y el INSERT fallaba con
    # «column "fecha" is of type date but expression is of type text».
    descartadas = 0
    cruce = conf.get("tambien_revisar")
    if cruce:
        otra, col_otra = cruce
        pos = columnas.index(conf["conflicto"])
        col_excel = indices[pos]
        vistos = ya_en_otra_tabla(
            otra, col_otra,
            [f[col_excel] if col_excel < len(f) else None for f in filas])
        if vistos:
            antes = len(filas)
            filas = [f for f in filas
                     if (f[col_excel] if col_excel < len(f) else None) not in vistos]
            descartadas = antes - len(filas)

    sql = (f'INSERT INTO plantacion.{tabla} ({destino}) VALUES %s '
           f'ON CONFLICT ("{conf["conflicto"]}") DO NOTHING RETURNING 1')

    nuevas = 0
    total = len(filas) + descartadas
    with db.get_cursor() as cur:
        for inicio in range(0, len(filas), LOTE):
            valores_fijos = list(fijos.values())
            bloque = []
            for f in filas[inicio:inicio + LOTE]:
                fila = []
                for n_col, i in enumerate(indices):
                    valor = f[i] if i < len(f) else None
                    if valor is None and rellenos[n_col] is not None:
                        valor = rellenos[n_col]
                    fila.append(valor)
                bloque.append(fila + valores_fijos)
            nuevas += len(execute_values(cur, sql, bloque, fetch=True))

    return nuevas, total - nuevas


def anotar(archivo: str, tabla: str, hash_archivo: str, leidas: int,
           nuevas: int, repetidas: int, ignoradas: list[str],
           resultado: str, detalle: str | None, usuario: str) -> None:
    """Deja constancia de la carga, salga bien o mal."""
    with db.get_cursor() as cur:
        cur.execute("""
            INSERT INTO plantacion.carga_archivos
                (archivo, tabla, huella, filas_leidas, filas_nuevas,
                 filas_repetidas, columnas_ignoradas, resultado, detalle, cargado_por)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (archivo[:200], tabla[:100], hash_archivo, leidas, nuevas, repetidas,
              ", ".join(ignoradas) or None, resultado,
              (detalle or "")[:2000] or None, usuario))


def historial(limite: int = 100) -> list[dict]:
    return db.fetch_all("""
        SELECT id, archivo, tabla, filas_leidas, filas_nuevas, filas_repetidas,
               columnas_ignoradas, resultado, detalle, cargado_por, cargado_at
        FROM plantacion.carga_archivos
        ORDER BY cargado_at DESC LIMIT %s
    """, (limite,))


def resumen_historial() -> dict:
    fila = db.fetch_one("""
        SELECT COUNT(*)                                       AS archivos,
               COUNT(*) FILTER (WHERE resultado = 'ok')       AS correctos,
               COUNT(*) FILTER (WHERE resultado <> 'ok')      AS con_error,
               COALESCE(SUM(filas_nuevas), 0)                 AS filas_nuevas,
               COALESCE(SUM(filas_repetidas), 0)              AS filas_repetidas,
               MAX(cargado_at)                                AS ultima
        FROM plantacion.carga_archivos
    """)
    return dict(fila) if fila else {}
